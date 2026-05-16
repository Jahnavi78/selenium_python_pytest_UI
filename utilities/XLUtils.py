import openpyxl

def getRowCount(file, sheetName):
    workbook= openpyxl.load_workbook(file)
    sheet=workbook[sheetName]
    return sheet.max_row

def getColumnCount(file, sheetName):
    workbook=openpyxl.load_workbook(file)
    sheet=workbook[sheetName]
    return sheet.max_column

def readData(file, sheetName, rownum, columnno):
    workbook=openpyxl.load_workbook(file)
    sheet=workbook[sheetName]
    return sheet.cell(row=rownum, column=columnno).value

def writeData(file, sheetName, rownum, columnno, data):
    workbook=openpyxl.load_workbook(file)
    sheet=workbook[sheetName]
    sheet.cell(row=rownum, column=columnno).value= data
    workbook.save(file)

def exceldatatodict(file, sheetName):
    workbook=openpyxl.load_workbook(file)
    sheet=workbook[sheetName]
    data={}
    for i in range(2, sheet.max_row+1):
        data[sheet.cell(row=i, column=1).value]=sheet.cell(row=i, column=2).value
    return data

